"""Relecture des fiches pédagogiques dans Excel, puis retour en JSONL.

    python -m cfdata.fiches_review export --out reports/relecture_fiches.xlsx
    (remplir les colonnes de notes dans Excel)
    python -m cfdata.fiches_review import --xlsx reports/relecture_fiches.xlsx --out data/annotated/fiches_relues.jsonl

Ou avec la page annotation/relecture_fiches.html (un fichier JSON de notes par relectrice) :
    python -m cfdata.fiches_review import-json --json annotation/relecture_ons.json,annotation/relecture_boutheina.json
"""
from __future__ import annotations

import argparse
import json
from datetime import date

from .common import read_jsonl, write_jsonl

CRIT = ["fidelite_objectifs", "donnees", "contraintes", "invention", "citations", "anonymisation", "faisabilite", "langue"]
HEAD = ["doc_id", "split", "langue", "document_anonymise", "objectifs", "donnees_fiche", "contraintes_fiche"] + CRIT + ["relu_par", "commentaire"]


def _txt(items):
    return "\n".join(f"- {i['texte']}  [{', '.join(i.get('sources', []))}]" for i in items)


def export(args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    fiches = read_jsonl("data/final/fiches_a_relire.jsonl")
    docs = {}
    for s in ("train", "dev", "test"):
        for d in read_jsonl(f"data/final/documents_{s}.jsonl"):
            docs[d["doc_id"]] = "\n\n".join(f"[{p['passage_id'].split('_')[-1]}] {p['text_anonymise']}" for p in d["passages"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Relecture"
    ws.append(HEAD)
    for f in fiches:
        ws.append([f["doc_id"], f.get("split"), f["lang"], docs.get(f["doc_id"], ""), _txt(f["objectifs"]),
                   _txt(f["donnees"]), _txt(f["contraintes"])] + [None] * len(CRIT) + [None, None])
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="2", showErrorMessage=True,
                        error="Note entre 0 et 2")
    ws.add_data_validation(dv)
    c0 = len(HEAD) - len(CRIT) - 2 + 1
    for col in range(c0, c0 + len(CRIT)):
        letter = ws.cell(1, col).column_letter
        dv.add(f"{letter}2:{letter}{len(fiches) + 1}")
        ws.column_dimensions[letter].width = 11
        for r in range(2, len(fiches) + 2):
            ws.cell(r, col).fill = PatternFill("solid", fgColor="FFF4CC")
    widths = {"A": 13, "B": 7, "C": 7, "D": 70, "E": 45, "F": 35, "G": 45}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "D2"
    g = wb.create_sheet("Grille")
    g.append(["Critère", "Question", "0 / 1 / 2"])
    for c, q in zip(CRIT, ["Chaque objectif découle-t-il du problème de la source ?",
                           "Les indicateurs et données utiles sont-ils tous repris ?",
                           "Toutes les contraintes sont présentes, sans en inventer ?",
                           "2 = aucune invention (chiffre, nom, contrainte, objectif absent de la source)",
                           "Chaque élément renvoie-t-il au bon passage [pXXX] ?",
                           "2 = aucun nom, site ou donnée sensible visible ni combinaison ré-identifiante",
                           "Consignes, livrables, critères adaptés à des étudiants ingénieurs ?",
                           "Langue correcte et homogène ?"]):
        g.append([c, q, "0 = absent ou faux · 1 = partiel · 2 = correct"])
    g.append([])
    g.append(["Acceptation", "total ≥ 13/16 et aucun 0 en invention ou anonymisation"])
    g.column_dimensions["A"].width = 20
    g.column_dimensions["B"].width = 80
    wb.save(args.out)
    print(f"{len(fiches)} fiches à relire -> {args.out}")


def imp(args):
    from openpyxl import load_workbook
    ws = load_workbook(args.xlsx).active
    rows = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(rows[0])}
    fiches = {f["doc_id"]: f for f in read_jsonl("data/final/fiches_a_relire.jsonl")}
    out, incomplete = [], 0
    for r in rows[1:]:
        scores = {c: r[idx[c]] for c in CRIT}
        if any(v is None for v in scores.values()) or not r[idx["relu_par"]]:
            incomplete += 1
            continue
        f = dict(fiches[r[idx["doc_id"]]])
        total = sum(int(v) for v in scores.values())
        f["validation"] = {"statut": "relue", "valide_par": r[idx["relu_par"]], "date": str(date.today()),
                           "scores": {k: int(v) for k, v in scores.items()}, "total": total,
                           "acceptee": total >= 13 and int(scores["invention"]) > 0 and int(scores["anonymisation"]) > 0,
                           "commentaire": r[idx["commentaire"]] or ""}
        out.append(f)
    write_jsonl(args.out, out)
    acc = sum(f["validation"]["acceptee"] for f in out)
    print(f"{len(out)} fiches relues ({acc} acceptées, {len(out) - acc} à corriger) ; {incomplete} lignes incomplètes ignorées -> {args.out}")


def imp_json(args):
    """Notes téléchargées depuis annotation/relecture_fiches.html -> fiches_relues.jsonl."""
    fiches = {f["doc_id"]: f for f in read_jsonl("data/final/fiches_a_relire.jsonl")}
    out, seen = [], set()
    for path in args.json.split(","):
        for r in json.load(open(path, encoding="utf-8")):
            scores = {c: r.get(c) for c in CRIT}
            if any(v not in (0, 1, 2) for v in scores.values()) or not r.get("relu_par"):
                raise SystemExit(f"{path} : notes incomplètes pour {r.get('doc_id')}")
            if r["doc_id"] in seen:
                raise SystemExit(f"{r['doc_id']} relue deux fois")
            seen.add(r["doc_id"])
            f = dict(fiches[r["doc_id"]])
            total = sum(scores.values())
            f["validation"] = {"statut": "relue", "valide_par": r["relu_par"], "date": args.date, "scores": scores,
                               "total": total, "acceptee": total >= 13 and scores["invention"] > 0 and scores["anonymisation"] > 0,
                               "commentaire": r.get("commentaire", ""), "fichier_notes": path.replace("\\", "/")}
            out.append(f)
    write_jsonl(args.out, out)
    acc = sum(f["validation"]["acceptee"] for f in out)
    print(f"{len(out)} fiches relues ({acc} acceptées, {len(out) - acc} à corriger) -> {args.out}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(required=True)
    a = sub.add_parser("export"); a.add_argument("--out", default="reports/relecture_fiches.xlsx"); a.set_defaults(f=export)
    b = sub.add_parser("import"); b.add_argument("--xlsx", default="reports/relecture_fiches.xlsx")
    b.add_argument("--out", default="data/annotated/fiches_relues.jsonl"); b.set_defaults(f=imp)
    c = sub.add_parser("import-json"); c.add_argument("--json", required=True)
    c.add_argument("--date", default=str(date.today()))
    c.add_argument("--out", default="data/annotated/fiches_relues.jsonl"); c.set_defaults(f=imp_json)
    args = ap.parse_args()
    args.f(args)


if __name__ == "__main__":
    main()
