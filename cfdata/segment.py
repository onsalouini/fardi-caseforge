"""SOURCES 2 et 3 — documents réels (écrits par l'équipe, publics, FARDI) -> passages + tâches Label Studio pré-annotées.

Usage :
    python -m cfdata.segment --raw data/raw --meta data/raw/metadata.csv --out data/interim

metadata.csv : doc_id,filename,source_id,source,secteur,licence,auteur,url
  source ∈ {equipe, public, fardi}
  source_id = identifiant de la source dans data/sources/registry.csv (ex. SRC-EQUIPE, SRC-PUB-003)

Sortie :
  data/interim/passages.jsonl          passages nettoyés avec leur provenance
  data/interim/labelstudio_tasks.json tâches à importer, avec pré-annotations (regex + lexiques) à corriger
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from pathlib import Path

from .common import detect_lang, log_step, write_jsonl
from .prelabel import ls_task

MIN_CHARS, MAX_CHARS = 25, 900


def read_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    if ext == ".docx":
        import docx  # python-docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # une ligne de tableau = un passage potentiel
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(parts)
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            text = "\n\n".join((page.extract_text() or "") for page in pdf.pages)
        if len(text.strip()) < 50:
            raise ValueError(f"{path} : PDF sans texte (scan). Passer d'abord par l'OCR : "
                             "ocrmypdf -l fra+ara+eng entree.pdf sortie.pdf")
        return text
    raise ValueError(f"Format non géré : {path}")


def clean(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace(" ", " ").replace("‏", "").replace("‎", "").replace("﻿", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"-\n(?=\w)", "", text)               # césure de fin de ligne
    text = re.sub(r"(?<![.\n:;])\n(?!\n)", " ", text)   # retour à la ligne au milieu d'une phrase
    return text.strip()


def split_passages(text: str) -> list[str]:
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    out: list[str] = []
    for b in blocks:
        if len(b) <= MAX_CHARS:
            out.append(b)
            continue
        cur = ""
        for sent in re.split(r"(?<=[.!?؟])\s+", b):
            if cur and len(cur) + len(sent) + 1 > MAX_CHARS:
                out.append(cur)
                cur = sent
            else:
                cur = f"{cur} {sent}".strip()
        if cur:
            out.append(cur)
    merged, carry = [], ""
    for p in out:  # fragments trop courts (titres, formules) fusionnés avec le passage suivant
        p = f"{carry} {p}".strip() if carry else p
        if len(p) < MIN_CHARS:
            carry = p
        else:
            merged.append(p)
            carry = ""
    if carry:
        if merged:
            merged[-1] += " " + carry
        else:
            merged.append(carry)
    return merged


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--meta", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    raw = Path(args.raw)
    with open(args.meta, encoding="utf-8") as f:
        meta = [m for m in csv.DictReader(f) if m.get("doc_id", "").strip()]

    passages, seen = [], set()
    for m in meta:
        doc_id = m["doc_id"].strip()
        if doc_id in seen:
            raise SystemExit(f"doc_id en double : {doc_id}")
        seen.add(doc_id)
        text = clean(read_document(raw / m["filename"]))
        parts = split_passages(text)
        doc_lang = max(("fr", "ar", "en"), key=lambda l: sum(detect_lang(p) == l for p in parts))
        for i, p in enumerate(parts):
            passages.append({
                "doc_id": doc_id, "passage_id": f"{doc_id}_p{i:03d}", "text": p,
                "lang_auto": detect_lang(p), "doc_lang": doc_lang,
                "source_id": m.get("source_id", ""), "source_type": m.get("source", ""),
                "secteur": m.get("secteur", ""), "licence": m.get("licence", ""),
                "url": m.get("url", ""), "auteur_document": m.get("auteur", ""),
                "famille": "FARDI" if m.get("source") == "fardi" else "R",
            })

    out = Path(args.out)
    n = write_jsonl(out / "passages.jsonl", passages)
    tasks = [ls_task(p) for p in passages]
    (out / "labelstudio_tasks.json").write_text(json.dumps(tasks, ensure_ascii=False, indent=1), encoding="utf-8")
    n_pre = sum(len(t["predictions"][0]["result"]) for t in tasks)
    log_step("segment", {"documents": len(meta), "passages": n, "pre_annotations": n_pre,
                         "sources": sorted({m.get("source_id", "") for m in meta})})
    print(f"{len(meta)} documents -> {n} passages ({n_pre} pré-annotations à vérifier)")
    print(f"  {out / 'passages.jsonl'}\n  {out / 'labelstudio_tasks.json'} (à importer dans Label Studio)")


if __name__ == "__main__":
    main()
